"""Generates demos/demo-vanilla/public/sample.xlsx, the demo's sample workbook.

    python scripts/sample-workbook.py          # needs openpyxl (3.1+)
    pwsh scripts/sample-workbook-excel.ps1     # optional: let Excel recalculate and re-save it,
                                               # and write tests/support/legacy-97.xls

A small but realistic sales and budget workbook that exercises what the viewer renders: three
sheets, formulas (including cross-sheet references), fonts, fills, borders, number formats, a
merged title, frozen panes, conditional formatting (colour scale, data bar, cell rule), a
drop-down data validation list, a tab colour and two charts.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / 'demos' / 'demo-vanilla' / 'public' / 'sample.xlsx'
GREEN = '1F9D63'
HEADER_FILL = PatternFill('solid', fgColor=GREEN)
BAND_FILL = PatternFill('solid', fgColor='E8F5EE')
HEADER_FONT = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
TITLE_FONT = Font(name='Calibri Light', size=18, bold=True, color='0E5A38')
THIN = Side(style='thin', color='B7C9BF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CURRENCY = '"$"#,##0;[Red]-"$"#,##0'
PERCENT = '0.0%'
MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
REGIONS = ['North', 'South', 'East', 'West', 'Central']
SALES = [
    [42100, 39800, 45250, 47100, 51230, 53900, 50110, 52480, 55120, 57900, 61200, 68450],
    [31200, 30150, 33900, 35020, 36880, 38110, 37450, 39200, 40310, 42780, 44900, 49800],
    [27800, 28900, 29450, 31200, 30980, 33450, 34120, 35800, 36020, 37950, 39400, 43150],
    [36500, 35100, 38750, 40200, 42150, 41980, 43300, 45120, 47800, 49010, 50500, 56300],
    [18900, 19450, 20100, 21850, 22400, 23900, 24150, 25020, 26300, 27100, 28850, 31900],
]


def header(ws, row, values, start_col=1):
    for offset, value in enumerate(values):
        cell = ws.cell(row=row, column=start_col + offset, value=value)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BOX
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def title(ws, text, span):
    ws.merge_cells(f'A1:{span}1')
    ws['A1'] = text
    ws['A1'].font = TITLE_FONT
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30


def sales_sheet(wb):
    ws = wb.active
    ws.title = 'Sales'
    ws.sheet_properties.tabColor = GREEN
    title(ws, 'Regional Sales 2026', 'O')
    ws['A2'] = 'Monthly revenue by region (USD). Totals and growth are formulas.'
    ws['A2'].font = Font(italic=True, color='5B6B63')
    header(ws, 3, ['Region', *MONTHS, 'Total', 'Growth'])
    for index, (region, values) in enumerate(zip(REGIONS, SALES)):
        row = 4 + index
        ws.cell(row=row, column=1, value=region).border = BOX
        for col, value in enumerate(values, start=2):
            cell = ws.cell(row=row, column=col, value=value)
            cell.number_format = CURRENCY
            cell.border = BOX
            if index % 2:
                cell.fill = BAND_FILL
        total = ws.cell(row=row, column=14, value=f'=SUM(B{row}:M{row})')
        total.number_format = CURRENCY
        total.font = Font(bold=True)
        total.border = BOX
        growth = ws.cell(row=row, column=15, value=f'=IF(B{row}=0,0,M{row}/B{row}-1)')
        growth.number_format = PERCENT
        growth.border = BOX
    total_row = 4 + len(REGIONS)
    ws.cell(row=total_row, column=1, value='Total').font = Font(bold=True)
    for col in range(2, 15):
        letter = ws.cell(row=3, column=col).column_letter
        cell = ws.cell(row=total_row, column=col, value=f'=SUM({letter}4:{letter}{total_row - 1})')
        cell.number_format = CURRENCY
        cell.font = Font(bold=True)
        cell.border = Border(top=Side(style='thin', color='0E5A38'), bottom=Side(style='double', color='0E5A38'))
    ws.cell(row=total_row + 2, column=1, value='Best month')
    ws.cell(row=total_row + 2, column=2, value=f'=INDEX(B3:M3,MATCH(MAX(B{total_row}:M{total_row}),B{total_row}:M{total_row},0))')
    ws.cell(row=total_row + 3, column=1, value='Average month')
    avg = ws.cell(row=total_row + 3, column=2, value=f'=AVERAGE(B{total_row}:M{total_row})')
    avg.number_format = CURRENCY
    ws.cell(row=total_row + 4, column=1, value='Report date')
    date = ws.cell(row=total_row + 4, column=2, value='=DATE(2026,12,31)')
    date.number_format = 'd-mmm-yyyy'

    ws.column_dimensions['A'].width = 14
    for col in range(2, 14):
        ws.column_dimensions[ws.cell(row=3, column=col).column_letter].width = 10.5
    ws.column_dimensions['N'].width = 13
    ws.column_dimensions['O'].width = 9
    ws.freeze_panes = 'B4'

    ws.conditional_formatting.add(
        f'B4:M{total_row - 1}',
        ColorScaleRule(start_type='min', start_color='F8696B', mid_type='percentile', mid_value=50,
                       mid_color='FFEB84', end_type='max', end_color='63BE7B'),
    )
    ws.conditional_formatting.add(
        f'N4:N{total_row - 1}',
        DataBarRule(start_type='min', end_type='max', color='1F9D63', showValue=True),
    )
    ws.conditional_formatting.add(
        f'O4:O{total_row - 1}',
        CellIsRule(operator='greaterThan', formula=['0.5'], font=Font(bold=True, color='006100'),
                   fill=PatternFill('solid', bgColor='C6EFCE')),
    )

    chart = BarChart()
    chart.type = 'col'
    chart.grouping = 'clustered'
    chart.title = 'Revenue by region'
    chart.y_axis.title = 'USD'
    chart.add_data(Reference(ws, min_col=14, min_row=3, max_row=total_row - 1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=4, max_row=total_row - 1))
    chart.height = 7.5
    chart.width = 16
    ws.add_chart(chart, f'A{total_row + 6}')
    return total_row


def budget_sheet(wb):
    ws = wb.create_sheet('Budget')
    ws.sheet_properties.tabColor = '2F75B5'
    title(ws, 'Operating Budget FY2026', 'F')
    header(ws, 3, ['Category', 'Owner', 'Budget', 'Actual', 'Variance', 'Status'])
    rows = [
        ('Salaries', 'People', 840000, 826500),
        ('Marketing', 'Growth', 145000, 162300),
        ('Cloud hosting', 'Engineering', 96000, 88750),
        ('Office and facilities', 'Operations', 72000, 74100),
        ('Travel', 'Sales', 38000, 29400),
        ('Training', 'People', 24000, 19850),
        ('Software licences', 'Engineering', 51000, 53600),
        ('Contingency', 'Finance', 30000, 0),
    ]
    for index, (category, owner, budget, actual) in enumerate(rows):
        row = 4 + index
        ws.cell(row=row, column=1, value=category)
        ws.cell(row=row, column=2, value=owner)
        ws.cell(row=row, column=3, value=budget).number_format = CURRENCY
        ws.cell(row=row, column=4, value=actual).number_format = CURRENCY
        ws.cell(row=row, column=5, value=f'=C{row}-D{row}').number_format = CURRENCY
        ws.cell(row=row, column=6, value=f'=IF(D{row}>C{row},"Over","On track")')
        for col in range(1, 7):
            ws.cell(row=row, column=col).border = BOX
    last = 3 + len(rows)
    total = last + 1
    ws.cell(row=total, column=1, value='Total').font = Font(bold=True)
    for col, letter in ((3, 'C'), (4, 'D'), (5, 'E')):
        cell = ws.cell(row=total, column=col, value=f'=SUM({letter}4:{letter}{last})')
        cell.number_format = CURRENCY
        cell.font = Font(bold=True)
    ws.cell(row=total + 1, column=1, value='Spent')
    ws.cell(row=total + 1, column=4, value=f'=D{total}/C{total}').number_format = PERCENT

    owners = DataValidation(type='list', formula1='"People,Growth,Engineering,Operations,Sales,Finance"',
                            allow_blank=False, showDropDown=False)
    owners.error = 'Pick an owner from the list.'
    owners.errorTitle = 'Unknown owner'
    owners.prompt = 'Choose the budget owner'
    owners.promptTitle = 'Owner'
    ws.add_data_validation(owners)
    owners.add(f'B4:B{last}')
    ws.conditional_formatting.add(
        f'E4:E{last}',
        CellIsRule(operator='lessThan', formula=['0'], font=Font(color='9C0006'),
                   fill=PatternFill('solid', bgColor='FFC7CE')),
    )
    for letter, width in zip('ABCDEF', (24, 14, 13, 13, 13, 11)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = 'A4'
    return total


def summary_sheet(wb, sales_total_row, budget_total_row):
    ws = wb.create_sheet('Summary')
    ws.sheet_properties.tabColor = 'C55A11'
    title(ws, 'Summary', 'D')
    header(ws, 3, ['Metric', 'Value', 'Target', 'Met'])
    metrics = [
        ('Annual revenue', f"=Sales!N{sales_total_row}", 2400000, CURRENCY),
        ('Operating cost', f"=Budget!D{budget_total_row}", 1300000, CURRENCY),
        ('Operating margin', '=(B4-B5)/B4', 0.4, PERCENT),
        ('Budget spent', f"=Budget!D{budget_total_row + 1}", 0.95, PERCENT),
        ('Regions over 500k', f'=COUNTIF(Sales!N4:N{sales_total_row - 1},">500000")', 3, '0'),
    ]
    for index, (label, formula, target, fmt) in enumerate(metrics):
        row = 4 + index
        ws.cell(row=row, column=1, value=label).border = BOX
        value = ws.cell(row=row, column=2, value=formula)
        value.number_format = fmt
        value.border = BOX
        goal = ws.cell(row=row, column=3, value=target)
        goal.number_format = fmt
        goal.border = BOX
        met = ws.cell(row=row, column=4, value=f'=IF(B{row}>=C{row},"Yes","No")')
        met.alignment = Alignment(horizontal='center')
        met.border = BOX
    ws['B6'].comment = Comment('Margin before tax; excludes one-off costs.', 'Finance')
    ws.cell(row=11, column=1, value='Month')
    ws.cell(row=11, column=2, value='All regions')
    for index, month in enumerate(MONTHS):
        row = 12 + index
        ws.cell(row=row, column=1, value=month)
        letter = ws.cell(row=3, column=2 + index).column_letter
        ws.cell(row=row, column=2, value=f'=Sales!{letter}{sales_total_row}').number_format = CURRENCY
    line = LineChart()
    line.title = 'Monthly revenue'
    line.add_data(Reference(ws, min_col=2, min_row=11, max_row=23), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=12, max_row=23))
    line.height = 7
    line.width = 14
    ws.add_chart(line, 'F3')
    for letter, width in zip('ABCD', (22, 15, 13, 8)):
        ws.column_dimensions[letter].width = width


def main():
    wb = Workbook()
    sales_total = sales_sheet(wb)
    budget_total = budget_sheet(wb)
    summary_sheet(wb, sales_total, budget_total)
    wb.properties.title = 'Sales and budget sample'
    wb.properties.creator = 'xlsx-viewer'
    wb.active = 0
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f'wrote {OUT}')


if __name__ == '__main__':
    main()
